"""Export the v3 document blocks as an editable Word document.

Heavy libraries are imported only by the explicit export request. There is no
OCR, model startup, network call, or temporary file in this module's lifecycle.
"""
from __future__ import annotations

import base64
from io import BytesIO
import math
import re
import struct
import xml.etree.ElementTree as ET

MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
CATEGORIES = ('选择题', '填空题', '判断题', '应用题')
NUMERALS = ('一', '二', '三', '四')
MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def _parts(value):
    """Read explicit math delimiters; malformed/bare TeX must not become text."""
    if not isinstance(value, str):
        raise ValueError('题目文字无效。')
    if len(value) > 80000:
        raise ValueError('题目文字过长。')
    cursor = 0
    while cursor < len(value):
        start = cursor
        while start < len(value):
            if value[start] == '$' and (start == 0 or value[start - 1] != '\\'):
                break
            start += 1
        prose = value[cursor:start]
        if re.search(r'\\[A-Za-z]+', prose):
            raise ValueError('题文含未标记的公式，请校对后再导出 Word。')
        if prose:
            yield False, prose.replace(r'\$', '$')
        if start == len(value):
            return
        delimiter = '$$' if value.startswith('$$', start) else '$'
        end = value.find(delimiter, start + len(delimiter))
        while end >= 0 and value[end - 1] == '\\':
            end = value.find(delimiter, end + len(delimiter))
        if end < 0:
            raise ValueError('公式标记不完整，请校对后再导出 Word。')
        expression = value[start + len(delimiter):end]
        if not expression.strip() or len(expression) > 20000:
            raise ValueError('公式为空或过长，请校对后再导出 Word。')
        yield True, expression
        cursor = end + len(delimiter)


def _native_math(expression):
    from docx.oxml import parse_xml
    from latex2mathml.converter import convert as to_mathml
    from mathml2omml import convert as to_omml

    try:
        mathml = to_mathml(expression)
        parsed = ET.fromstring(mathml)
        if any(re.search(r'\\[A-Za-z]+', node.text or '') for node in parsed.iter()):
            raise ValueError('Unrecognised TeX command')
        converted = to_omml(mathml)
        wrapper = parse_xml(f'<root xmlns:m="{MATH_NS}">{converted}</root>')
        nodes = [node for node in wrapper.iter() if node.tag == f'{{{MATH_NS}}}oMath']
        if len(nodes) != 1 or re.search(r'\\[A-Za-z]+', ''.join(nodes[0].itertext())):
            raise ValueError('No complete native equation')
        return nodes[0]
    except Exception as error:
        raise ValueError('有公式无法转换为可编辑的 Word 公式，请校对后重试；未生成包含公式源码的文档。') from error


def _rich(paragraph, value):
    for is_math, content in _parts(value):
        if is_math:
            paragraph._p.append(_native_math(content))
        else:
            paragraph.add_run(content)


def _answer_space(question):
    value = question.get('answerSpaceMm', 24 if question.get('category') == '应用题' else 0)
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 <= value <= 96:
        raise ValueError('答题留白须为 0 至 96 毫米。')
    return value


def _image(group):
    if not isinstance(group, dict):
        raise ValueError('图形尚未准备完成，请校对后再导出。')
    value = group.get('dataBase64')
    if not isinstance(value, str) or len(value) > 28 * 1024 * 1024:
        raise ValueError('图形图片数据无效或过大。')
    if value.startswith('data:image/png;base64,'):
        value = value.split(',', 1)[1]
    try:
        data = base64.b64decode(value, validate=True)
    except Exception as error:
        raise ValueError('图形图片数据无效。') from error
    if len(data) < 24 or data[:8] != b'\x89PNG\r\n\x1a\n':
        raise ValueError('图形必须是已校对的 PNG 图片。')
    width, height = struct.unpack('!II', data[16:24])
    if not (0 < width <= 20000 and 0 < height <= 20000):
        raise ValueError('图形图片尺寸无效。')
    for key, actual in (('widthPx', width), ('heightPx', height)):
        declared = group.get(key)
        if isinstance(declared, bool) or not isinstance(declared, (int, float)) or not math.isfinite(declared) or abs(declared - actual) > 1:
            raise ValueError('图形图片尺寸与排版数据不一致。')
    width_mm, height_mm = 170, 170 * height / width
    if height_mm > 230:
        width_mm *= 230 / height_mm
        height_mm = 230
    return data, width_mm, height_mm


def _table(document, block):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm

    columns, items = block.get('columns', 1), block.get('items')
    if not isinstance(columns, int) or isinstance(columns, bool) or not 1 <= columns <= 4 or not isinstance(items, list) or len(items) > 12:
        raise ValueError('选项的列数或内容无效。')
    if not items:
        return []
    table = document.add_table(rows=math.ceil(len(items) / columns), cols=columns)
    table.autofit = False
    properties = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'nil')
        borders.append(border)
    properties.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        item = OxmlElement('w:' + side)
        item.set(qn('w:w'), '0' if side != 'right' else '57')
        item.set(qn('w:type'), 'dxa')
        margins.append(item)
    properties.append(margins)
    paragraphs = []
    for index, cell in enumerate(cell for row in table.rows for cell in row.cells):
        cell.width = Mm(170 / columns)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Mm(0)
        paragraphs.append(p)
        if index < len(items):
            item = items[index]
            if not isinstance(item, dict) or not isinstance(item.get('label'), str):
                raise ValueError('选项标签无效。')
            p.add_run(item['label'] + '. ')
            _rich(p, item.get('text', ''))
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    return paragraphs


def export_word(params):
    """Return an in-memory DOCX using the exact order of the canonical blocks."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    if not isinstance(params, dict) or not isinstance(params.get('project'), dict):
        raise ValueError('试卷工程无效。')
    project = params['project']
    questions = project.get('questions')
    if project.get('version') != 3 or not isinstance(questions, list) or not questions or len(questions) > 200:
        raise ValueError('请先添加并校对题目，再导出 Word。')
    if any(not isinstance(q, dict) or q.get('category') not in CATEGORIES for q in questions):
        raise ValueError('题目分类无效。')
    groups = params.get('figureGroups', {})
    if not isinstance(groups, dict):
        raise ValueError('图形排版数据无效。')
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(20)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = '宋体', Pt(12)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.line_spacing = 1.6
    normal.paragraph_format.space_after = Mm(1)
    title = project.get('title', '数学练习')
    if not isinstance(title, str) or len(title) > 300:
        raise ValueError('试卷标题无效。')
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    title_run.bold, title_run.font.size = True, Pt(19)
    title_paragraph.paragraph_format.space_after = Mm(7)
    title_paragraph.paragraph_format.keep_with_next = True
    category_index = 0
    for category in CATEGORIES:
        selected = [q for q in questions if q['category'] == category]
        if not selected:
            continue
        category_paragraph = document.add_paragraph()
        category_run = category_paragraph.add_run(f'{NUMERALS[category_index]}、{category}')
        category_run.bold, category_run.font.size = True, Pt(13)
        category_paragraph.paragraph_format.space_before = Mm(4 if category_index else 0)
        category_paragraph.paragraph_format.space_after = Mm(2)
        category_paragraph.paragraph_format.keep_with_next = True
        category_index += 1
        for number, question in enumerate(selected, 1):
            blocks = question.get('blocks')
            if not isinstance(blocks, list) or not blocks:
                raise ValueError('题目内容块为空，请先完成校对。')
            paragraphs, numbered = [], False
            for block in blocks:
                if not isinstance(block, dict):
                    raise ValueError('题目内容块无效。')
                kind = block.get('type')
                if kind == 'text':
                    p = document.add_paragraph()
                    if not numbered:
                        p.add_run(f'{number}、')
                        p.paragraph_format.space_before = Mm(2 if number > 1 else 0)
                        numbered = True
                    if block.get('label'):
                        if not isinstance(block['label'], str):
                            raise ValueError('小问标签无效。')
                        p.add_run(block['label'] + ' ')
                    _rich(p, block.get('text', ''))
                    paragraphs.append(p)
                elif kind == 'options':
                    paragraphs.extend(_table(document, block))
                elif kind == 'figures':
                    if not block.get('figureIds'):
                        continue
                    image, width, height = _image(groups.get(block.get('id')))
                    p = document.add_paragraph()
                    p.paragraph_format.space_before = Mm(1)
                    p.paragraph_format.space_after = Mm(1)
                    shape = p.add_run().add_picture(BytesIO(image), width=Mm(width), height=Mm(height))
                    captions = [str(f.get('caption', '')) for f in question.get('figures', []) if f.get('id') in block['figureIds']]
                    shape._inline.docPr.set('descr', '、'.join(c for c in captions if c and not re.fullmatch(r'(?:题图|提图|第\d+题图)', c)) or '几何图形')
                    paragraphs.append(p)
                else:
                    raise ValueError('题目含不支持的内容块。')
            if not numbered:
                raise ValueError('题目缺少题文，请完成校对后再导出。')
            blank = _answer_space(question)
            if blank:
                p = document.add_paragraph()
                p.add_run(' ').font.size = Pt(1)
                p.paragraph_format.line_spacing = Mm(blank)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.space_before = p.paragraph_format.space_after = Mm(0)
                paragraphs.append(p)
            for index, p in enumerate(paragraphs):
                p.paragraph_format.keep_together = True
                p.paragraph_format.keep_with_next = index < len(paragraphs) - 1
    buffer = BytesIO()
    document.save(buffer)
    filename = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', title).strip(' .')[:100] or '数学练习'
    return {'fileName': filename + '.docx', 'mimeType': MIME, 'dataBase64': base64.b64encode(buffer.getvalue()).decode('ascii')}
